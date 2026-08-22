"""Validate and apply the owner-managed Wildlife content workbook.

The workbook intentionally controls only Wildlife's collection projection and game values.
iNaturalist taxon identity stays in catalogues/taxa.yaml and must already exist before a row can
be added to a regional catalogue.
"""

from __future__ import annotations

import argparse
import json
import os
import tempfile
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
CATALOGUES = ROOT / "catalogues"
PILOTS = ("mediterranean_europe", "east_africa", "caribbean")
RARITIES = {"unknown", "common", "uncommon", "rare", "very_rare"}
PRESTIGE = {"standard", "legendary"}
EVENT_KEYS = (
    "confirmed_observation", "first_species", "research_grade", "regional_discovery",
    "regional_legend", "regional_essentials", "regional_icons", "identification_given",
    "anomaly_confirmed",
)


def fail(message: str) -> None:
    raise ValueError(f"Workbook validation failed: {message}")


def write_json_atomically(path: Path, data: object) -> None:
    payload = json.dumps(data, indent=2, ensure_ascii=False) + "\n"
    with tempfile.NamedTemporaryFile("w", encoding="utf-8", delete=False, dir=path.parent) as temp:
        temp.write(payload)
        temp_path = Path(temp.name)
    os.replace(temp_path, path)


def rows(sheet):
    return list(sheet.iter_rows(min_row=2, values_only=True))


def nonempty(row: tuple[object, ...]) -> bool:
    return any(value not in (None, "") for value in row)


def normalise_bool(value: object, label: str) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, str) and value.strip().lower() in {"true", "false"}:
        return value.strip().lower() == "true"
    fail(f"{label} must be TRUE or FALSE")


def positive_int(value: object, label: str, zero_allowed: bool = False) -> int:
    if isinstance(value, bool) or not isinstance(value, (int, float)) or int(value) != value:
        fail(f"{label} must be a whole number")
    value = int(value)
    if value < 0 or (not zero_allowed and value == 0):
        fail(f"{label} must be {'zero or greater' if zero_allowed else 'positive'}")
    return value


def load_workbook_content(workbook_path: Path) -> tuple[dict, dict, dict]:
    try:
        workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    except FileNotFoundError:
        fail(f"missing workbook: {workbook_path}")
    required_sheets = {"Catalogue", "Checklists", "XP and levels"}
    missing = required_sheets.difference(workbook.sheetnames)
    if missing:
        fail(f"missing sheet(s): {', '.join(sorted(missing))}")

    global_taxa = {
        entry["taxon_id"]: entry
        for entry in json.loads((CATALOGUES / "taxa.yaml").read_text(encoding="utf-8"))["taxa"]
    }
    catalogue_by_region: dict[str, dict] = {}
    seen_taxa: dict[str, set[int]] = {region: set() for region in PILOTS}
    for index, row in enumerate(rows(workbook["Catalogue"]), start=2):
        if not nonempty(row):
            continue
        region, version, taxon_id, _, _, rarity, prestige, seasonality, provenance = row[:9]
        if region not in PILOTS:
            fail(f"Catalogue row {index} has an unsupported region")
        if not isinstance(version, str) or not version.strip():
            fail(f"Catalogue row {index} needs a catalogue version")
        taxon_id = positive_int(taxon_id, f"Catalogue row {index} taxon ID")
        if taxon_id not in global_taxa:
            fail(f"Catalogue row {index} taxon {taxon_id} is not in catalogues/taxa.yaml")
        if taxon_id in seen_taxa[region]:
            fail(f"Catalogue row {index} duplicates taxon {taxon_id} in {region}")
        seen_taxa[region].add(taxon_id)
        rarity = str(rarity or "").strip()
        prestige = str(prestige or "").strip()
        if rarity not in RARITIES or prestige not in PRESTIGE:
            fail(f"Catalogue row {index} needs a supported encounter rarity and prestige")
        if not isinstance(seasonality, str) or not seasonality.strip() or not isinstance(provenance, str) or not provenance.strip():
            fail(f"Catalogue row {index} needs seasonality and inclusion provenance")
        destination = catalogue_by_region.setdefault(region, {"version": version.strip(), "taxa": []})
        if destination["version"] != version.strip():
            fail(f"Catalogue row {index} conflicts with the {region} catalogue version")
        destination["taxa"].append({
            "taxon_id": taxon_id,
            "encounter_rarity": rarity,
            "prestige": prestige,
            "seasonality": {"en": seasonality.strip()},
            "inclusion_provenance": provenance.strip(),
        })
    if set(catalogue_by_region) != set(PILOTS):
        fail("Catalogue must retain at least one entry for every pilot region")

    checklists: dict[str, dict[str, list[tuple[int, int]]]] = {
        region: {"essential": [], "icon": []} for region in PILOTS
    }
    for index, row in enumerate(rows(workbook["Checklists"]), start=2):
        if not nonempty(row):
            continue
        region, kind, position, taxon_id = row[:4]
        if region not in PILOTS or kind not in {"essential", "icon"}:
            fail(f"Checklist row {index} has an unsupported region or list")
        position = positive_int(position, f"Checklist row {index} position")
        taxon_id = positive_int(taxon_id, f"Checklist row {index} taxon ID")
        if taxon_id not in seen_taxa[region]:
            fail(f"Checklist row {index} taxon {taxon_id} is not in the {region} catalogue")
        checklists[region][kind].append((position, taxon_id))
    achievements: dict[str, dict] = {}
    for region, lists in checklists.items():
        expected = {"essential": 10, "icon": 5}
        resolved: dict[str, list[int]] = {}
        for kind, count in expected.items():
            entries = sorted(lists[kind])
            if len(entries) != count:
                fail(f"{region} must have exactly {count} {kind} rows")
            positions = [position for position, _ in entries]
            taxon_ids = [taxon_id for _, taxon_id in entries]
            if positions != list(range(1, count + 1)) or len(set(taxon_ids)) != count:
                fail(f"{region} {kind} rows need unique positions 1–{count} and taxa")
            resolved[kind] = taxon_ids
        if set(resolved["essential"]).intersection(resolved["icon"]):
            fail(f"{region} cannot use a taxon in both Essentials and Icons")
        for entry in catalogue_by_region[region]["taxa"]:
            if entry["taxon_id"] in resolved["essential"]:
                entry["prestige"] = "standard"
            elif entry["taxon_id"] in resolved["icon"]:
                entry["prestige"] = "legendary"
        achievements[region] = {
            "region": region,
            "catalogue_version": catalogue_by_region[region]["version"],
            "essentials": resolved["essential"],
            "icons": resolved["icon"],
        }

    config_rows = [row for row in rows(workbook["XP and levels"]) if nonempty(row)]
    values = {str(row[0]).strip(): row[1:3] for row in config_rows if row[0] not in (None, "")}
    if "rules_version" not in values or not isinstance(values["rules_version"][0], str):
        fail("XP and levels needs a rules_version")
    missing_events = set(EVENT_KEYS).difference(values)
    if missing_events:
        fail(f"XP and levels is missing {', '.join(sorted(missing_events))}")
    events = {
        key: {"xp": positive_int(values[key][0], f"XP value for {key}", zero_allowed=True), "enabled": normalise_bool(values[key][1], f"Enabled flag for {key}")}
        for key in EVENT_KEYS
    }
    rarity_xp = {}
    for rarity in ("common", "uncommon", "rare", "very_rare"):
        key = f"rarity_{rarity}"
        if key not in values:
            fail(f"XP and levels is missing {key}")
        rarity_xp[rarity] = positive_int(values[key][0], f"XP value for {key}", zero_allowed=True)

    levels = []
    for index, row in enumerate(workbook["XP and levels"].iter_rows(min_row=2, min_col=6, max_col=8, values_only=True), start=2):
        if not nonempty(row):
            continue
        key, display_name, threshold = row
        if not isinstance(key, str) or not key.strip() or not isinstance(display_name, str) or not display_name.strip():
            fail(f"Level row {index} needs a key and display name")
        levels.append({"key": key.strip(), "display_name": display_name.strip(), "threshold_xp": positive_int(threshold, f"Level row {index} threshold", zero_allowed=True)})
    if not levels or levels[0]["threshold_xp"] != 0 or len({level["key"] for level in levels}) != len(levels):
        fail("Levels need unique keys and must start at 0 XP")
    if any(later["threshold_xp"] <= earlier["threshold_xp"] for earlier, later in zip(levels, levels[1:])):
        fail("Level thresholds must strictly increase")

    repeats = []
    for index, row in enumerate(workbook["XP and levels"].iter_rows(min_row=2, min_col=10, max_col=12, values_only=True), start=2):
        if not nonempty(row):
            continue
        event, occurrence, points = row
        if event != "confirmed_observation":
            fail(f"Repeat observation row {index} has an unsupported event")
        repeats.append((positive_int(occurrence, f"Repeat row {index} occurrence"), positive_int(points, f"Repeat row {index} XP", zero_allowed=True)))
    repeats = sorted(repeats)
    if not repeats or [occurrence for occurrence, _ in repeats] != list(range(1, len(repeats) + 1)):
        fail("Repeat observation occurrences must start at 1 and have no gaps")

    progression = {
        "schema_version": 1,
        "version": values["rules_version"][0].strip(),
        "events": events,
        "rarity_xp": rarity_xp,
        "repeat_observation_xp": [points for _, points in repeats],
        "levels": levels,
    }
    return catalogue_by_region, achievements, progression


def kotlin_config(progression: dict) -> str:
    events = progression["events"]
    rarity = progression["rarity_xp"]
    level_lines = "\n".join(
        f'        ProgressionLevel("{level["key"]}", "{level["display_name"]}", {level["threshold_xp"]}),'
        for level in progression["levels"]
    )
    return f'''// Generated by tools/import_content_workbook.py. Do not edit by hand.\npackage com.wildlife.feasibility\n\ninternal object GeneratedProgressionConfig {{\n    const val VERSION = "{progression["version"]}"\n    val confirmedObservationXp = {events["confirmed_observation"]["xp"]}\n    val firstSpeciesXp = {events["first_species"]["xp"]}\n    val researchGradeXp = {events["research_grade"]["xp"]}\n    val researchGradeEnabled = {str(events["research_grade"]["enabled"]).lower()}\n    val regionalDiscoveryXp = {events["regional_discovery"]["xp"]}\n    val regionalLegendXp = {events["regional_legend"]["xp"]}\n    val regionalEssentialsXp = {events["regional_essentials"]["xp"]}\n    val regionalIconsXp = {events["regional_icons"]["xp"]}\n    val identificationGivenXp = {events["identification_given"]["xp"]}\n    val identificationGivenEnabled = {str(events["identification_given"]["enabled"]).lower()}\n    val anomalyConfirmedXp = {events["anomaly_confirmed"]["xp"]}\n    val anomalyConfirmedEnabled = {str(events["anomaly_confirmed"]["enabled"]).lower()}\n    val rarityXp = mapOf(\n        EncounterRarity.COMMON to {rarity["common"]},\n        EncounterRarity.UNCOMMON to {rarity["uncommon"]},\n        EncounterRarity.RARE to {rarity["rare"]},\n        EncounterRarity.VERY_RARE to {rarity["very_rare"]},\n        EncounterRarity.UNKNOWN to 0,\n    )\n    val repeatObservationXp = listOf({", ".join(str(value) for value in progression["repeat_observation_xp"])})\n    val levels = listOf(\n{level_lines}\n    )\n}}\n'''


def apply(catalogues: dict, achievements: dict, progression: dict) -> None:
    updates: list[tuple[Path, object]] = []
    for region in PILOTS:
        path = CATALOGUES / "regions" / region / "catalogue.yaml"
        current = json.loads(path.read_text(encoding="utf-8"))
        current["catalogue_version"] = catalogues[region]["version"]
        current["taxa"] = catalogues[region]["taxa"]
        updates.append((path, current))
    updates.append((CATALOGUES / "achievements.yaml", {"schema_version": 1, "achievements": [achievements[region] for region in PILOTS]}))
    updates.append((CATALOGUES / "progression.yaml", progression))
    for path, data in updates:
        write_json_atomically(path, data)
    generated = ROOT / "app" / "src" / "main" / "java" / "com" / "ecotracker" / "feasibility" / "GeneratedProgressionConfig.kt"
    generated.write_text(kotlin_config(progression), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=CATALOGUES / "review" / "wildlife-content-manager.xlsx")
    parser.add_argument("--apply", action="store_true", help="Write the validated workbook into source files.")
    args = parser.parse_args()
    catalogues, achievements, progression = load_workbook_content(args.workbook)
    if args.apply:
        apply(catalogues, achievements, progression)
        print("Applied validated content workbook to catalogue and progression sources.")
    else:
        print("Content workbook is valid. Re-run with --apply to update sources.")


if __name__ == "__main__":
    main()
